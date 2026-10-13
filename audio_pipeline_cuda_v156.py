#!/usr/bin/env python3
"""Audio2Text 1.5.6: local multilingual ASR, lossless source and resumable translation.

Import has no filesystem/model side effects. v154 and its cache are never imported.
Generated translations remain needs_review until human semantic validation; technical
coverage is reported separately from semantic completeness.
"""

from __future__ import annotations

import argparse
import copy
import gc
import hashlib
import importlib.metadata
import json
import logging
import math
import os
import re
import shutil
import subprocess
import sys
import time
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Callable, Protocol

ROOT = Path(__file__).resolve().parent
VERSION = "1.5.6"
SCHEMA = 2
BUILD = "v156-afnl-coverage-1"
LOG = logging.getLogger("audio2text")
ASR_LANGUAGES = {"ru", "en", "zh", "af", "it", "he", "de", "fr", "nl"}
NLLB = {
    "ru": "rus_Cyrl",
    "en": "eng_Latn",
    "zh": "zho_Hans",
    "af": "afr_Latn",
    "it": "ita_Latn",
    "he": "heb_Hebr",
    "de": "deu_Latn",
    "fr": "fra_Latn",
    "nl": "nld_Latn",
    "zu": "zul_Latn",
}
ROLES = {
    "ДОКЛАДЧИК": "007058",
    "ПЕРЕВОДЧИК": "1F4E79",
    "УЧАСТНИК": "802869",
    "НЕОПРЕДЕЛЁННЫЙ ГОЛОС": "666666",
}
TRANSLATION_STATUSES = {
    "pending",
    "complete",
    "partial",
    "needs_review",
    "not_required",
    "disabled",
    "unsupported_language",
    "failed",
}
DEFAULTS = {
    "INPUT": "input",
    "OUTPUT": "output",
    "TEMP": "temp/v156",
    "WHISPER_MODEL": "large-v3",
    "TRANSLATOR_MODEL": "facebook/nllb-200-distilled-600M",
    "WHISPER_PATH": None,
    "TRANSLATOR_PATH": None,
    "DIARIZATION_PATH": None,
    "FORCE_LANGUAGE": None,
    "AUTO_LANGUAGE": True,
    "LANGUAGE_CONFIDENCE": 0.50,
    "LANGUAGE_PROBE_SECONDS": 6.0,
    "ASR_UNIT_SECONDS": 24.0,
    "BOUNDARY_PADDING_SECONDS": 1.0,
    "MIN_LANGUAGE_UNIT_SECONDS": 3.0,
    "USE_VAD": False,
    "VAD_COMPARE": True,
    "VAD_THRESHOLD": 0.65,
    "VAD_MIN_SPEECH_MS": 250,
    "VAD_MIN_SILENCE_MS": 600,
    "VAD_PAD_MS": 180,
    "BEAM_SIZE": 5,
    "BEST_OF": 5,
    "TEMPERATURE": 0.0,
    "CONDITION_ON_PREVIOUS_TEXT": False,
    "REPETITION_PENALTY": 1.0,
    "NO_REPEAT_NGRAM_SIZE": 0,
    "NO_SPEECH_THRESHOLD": 0.65,
    "LOG_PROB_THRESHOLD": -1.0,
    "COMPRESSION_RATIO_THRESHOLD": 2.3,
    "MIN_SPEAKERS": 1,
    "MAX_SPEAKERS": None,
    "SAMPLE_RATE": 16000,
    "DEVICE": "auto",
    "GPU_FP16": True,
    "TRANSLATION_DEVICE": "cpu",
    "CPU_THREADS": 8,
    "TRANSLATE_MODE": "auto",
    "TRANSLATION_TOKEN_BUDGET": 192,
    "TRANSLATION_MAX_NEW_TOKENS": 384,
    "TRANSLATION_RETRY_DEPTH": 2,
    "TRANSLATION_BEAMS": 4,
    "CHINESE_SCRIPT": "auto",
    "JSON_SAVE_INTERVAL": 1,
    "TCM_DICT": {},
    "TERMS": {},
    "DICTIONARY_FILES": [],
    "SPEAKER_ROLES": {},
    "SUPPORTED_EXTENSIONS": [
        ".mp3",
        ".wav",
        ".flac",
        ".aac",
        ".m4a",
        ".ogg",
        ".opus",
        ".mp4",
        ".mkv",
        ".avi",
        ".mov",
        ".webm",
    ],
}


class PipelineError(RuntimeError):
    """An actionable pipeline failure (no credentials in messages)."""


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def digest(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(
            value,
            sort_keys=True,
            ensure_ascii=False,
            allow_nan=False,
            separators=(",", ":"),
        ).encode()
    ).hexdigest()


def text_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(4 * 1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def read_json(path: Path) -> Any:
    try:
        return json.loads(
            path.read_text(encoding="utf-8-sig"),
            parse_constant=lambda s: (_ for _ in ()).throw(ValueError(s)),
        )
    except json.JSONDecodeError as error:
        raise PipelineError(
            f"JSON {path.name}: line {error.lineno}, column {error.colno}: {error.msg}"
        ) from error
    except (OSError, UnicodeError, ValueError) as error:
        raise PipelineError(
            f"Не удалось прочитать JSON {path.name}: {type(error).__name__}"
        ) from error


def validate_config(config: dict) -> dict:
    c = copy.deepcopy(config)
    for key in (
        "AUTO_LANGUAGE",
        "USE_VAD",
        "VAD_COMPARE",
        "GPU_FP16",
        "CONDITION_ON_PREVIOUS_TEXT",
    ):
        if type(c[key]) is not bool:
            raise PipelineError(f"Config {key}: требуется boolean")
    bounds = {
        "BEAM_SIZE": (1, 20),
        "BEST_OF": (1, 20),
        "MIN_SPEAKERS": (1, 100),
        "JSON_SAVE_INTERVAL": (1, 1000),
        "TRANSLATION_TOKEN_BUDGET": (8, 512),
        "TRANSLATION_MAX_NEW_TOKENS": (8, 1024),
        "TRANSLATION_RETRY_DEPTH": (0, 4),
        "TRANSLATION_BEAMS": (1, 10),
        "NO_REPEAT_NGRAM_SIZE": (0, 20),
        "VAD_MIN_SPEECH_MS": (0, 5000),
        "VAD_MIN_SILENCE_MS": (0, 5000),
        "VAD_PAD_MS": (0, 2000),
        "CPU_THREADS": (1, 64),
    }
    for key, (low, high) in bounds.items():
        if type(c[key]) is not int or not low <= c[key] <= high:
            raise PipelineError(f"Config {key}: integer {low}..{high}")
    if c["MAX_SPEAKERS"] is not None and (
        type(c["MAX_SPEAKERS"]) is not int
        or not c["MIN_SPEAKERS"] <= c["MAX_SPEAKERS"] <= 100
    ):
        raise PipelineError(
            "Config MAX_SPEAKERS: null или integer >= MIN_SPEAKERS, <=100"
        )
    if type(c["SAMPLE_RATE"]) is not int or c["SAMPLE_RATE"] != 16000:
        raise PipelineError("SAMPLE_RATE должен быть 16000 для Whisper ndarray")
    for key, low, high in [
        ("LANGUAGE_CONFIDENCE", 0, 1),
        ("VAD_THRESHOLD", 0, 1),
        ("NO_SPEECH_THRESHOLD", 0, 1),
        ("LANGUAGE_PROBE_SECONDS", 1, 30),
        ("ASR_UNIT_SECONDS", 3, 30),
        ("MIN_LANGUAGE_UNIT_SECONDS", 1, 10),
        ("BOUNDARY_PADDING_SECONDS", 0, 2),
        ("TEMPERATURE", 0, 1),
        ("REPETITION_PENALTY", 1, 2),
        ("LOG_PROB_THRESHOLD", -10, 0),
        ("COMPRESSION_RATIO_THRESHOLD", 1, 10),
    ]:
        v = c[key]
        if (
            isinstance(v, bool)
            or not isinstance(v, (int, float))
            or not math.isfinite(v)
            or not low <= v <= high
        ):
            raise PipelineError(f"Config {key}: number {low}..{high}")
    if c["MIN_LANGUAGE_UNIT_SECONDS"] > c["ASR_UNIT_SECONDS"]:
        raise PipelineError("MIN_LANGUAGE_UNIT_SECONDS > ASR_UNIT_SECONDS")
    for key, choices in [
        ("DEVICE", {"auto", "cpu", "cuda"}),
        ("TRANSLATION_DEVICE", {"cpu", "cuda"}),
        ("TRANSLATE_MODE", {"auto", "off"}),
        ("CHINESE_SCRIPT", {"auto", "Hans", "Hant"}),
    ]:
        if c[key] not in choices:
            raise PipelineError(f"Config {key}: допустимы {sorted(choices)}")
    force = c["FORCE_LANGUAGE"]
    if force is not None:
        if not isinstance(force, str):
            raise PipelineError("FORCE_LANGUAGE: string или null")
        force = {"iw": "he"}.get(force.strip().lower(), force.strip().lower())
        if force not in ASR_LANGUAGES | {"zu"}:
            raise PipelineError(f"Unsupported FORCE_LANGUAGE: {force!r}")
        c["FORCE_LANGUAGE"] = force
    if not c["AUTO_LANGUAGE"] and not force:
        raise PipelineError("AUTO_LANGUAGE=false требует явный FORCE_LANGUAGE")
    for key in ("TCM_DICT", "TERMS"):
        if not isinstance(c[key], dict) or not all(
            isinstance(k, str) and k and isinstance(v, str) for k, v in c[key].items()
        ):
            raise PipelineError(
                f"Config {key}: требуется object непустая строка -> строка"
            )
    if not isinstance(c["DICTIONARY_FILES"], list) or not all(
        isinstance(x, str) for x in c["DICTIONARY_FILES"]
    ):
        raise PipelineError("DICTIONARY_FILES: список путей")
    if not isinstance(c["SPEAKER_ROLES"], dict):
        raise PipelineError("SPEAKER_ROLES: source SHA-256 -> speaker mapping")
    for source_id, roles in c["SPEAKER_ROLES"].items():
        if (
            not re.fullmatch(r"[a-f0-9]{64}:[a-f0-9]{64}", source_id)
            or not isinstance(roles, dict)
            or not all(isinstance(k, str) and v in ROLES for k, v in roles.items())
        ):
            raise PipelineError(
                "SPEAKER_ROLES: ключ source_sha256:processing_fingerprint; проверьте роли"
            )
    for key in ("INPUT", "OUTPUT", "TEMP", "WHISPER_MODEL", "TRANSLATOR_MODEL"):
        if not isinstance(c[key], str) or not c[key].strip():
            raise PipelineError(f"Config {key}: непустая строка")
    for key in ("WHISPER_PATH", "TRANSLATOR_PATH", "DIARIZATION_PATH"):
        if c[key] is not None and not isinstance(c[key], str):
            raise PipelineError(f"Config {key}: path string или null")
    if c["WHISPER_MODEL"] != "large-v3":
        raise PipelineError(
            "В этой версии проверен WHISPER_MODEL=large-v3; другой backend требует отдельной проверки"
        )
    if not isinstance(c["SUPPORTED_EXTENSIONS"], list) or not all(
        isinstance(x, str) and x.startswith(".") for x in c["SUPPORTED_EXTENSIONS"]
    ):
        raise PipelineError("SUPPORTED_EXTENSIONS: список расширений")
    return c


def load_config(path: Path | None = None) -> dict:
    c = copy.deepcopy(DEFAULTS)
    # Import only existing terminology, never legacy destructive cleanup or global settings.
    legacy = ROOT / "config.json"
    if legacy.exists():
        data = read_json(legacy)
        if not isinstance(data, dict):
            raise PipelineError("config.json: ожидается object")
        for key in ("TCM_DICT", "TERMS"):
            if key in data:
                c[key] = data[key]
    path = path or ROOT / "config_v156.json"
    if path.exists():
        data = read_json(path)
        if not isinstance(data, dict):
            raise PipelineError(f"{path.name}: ожидается object")
        unknown = sorted(set(data) - set(DEFAULTS))
        if unknown:
            raise PipelineError(
                f"{path.name}: неизвестные настройки: {', '.join(unknown)}"
            )
        c.update(data)
    c = validate_config(c)
    extras = []
    for filename in c["DICTIONARY_FILES"]:
        p = inside_root(filename)
        data = read_json(p)
        if not isinstance(data, dict) or not all(
            isinstance(k, str) and k and isinstance(v, str) for k, v in data.items()
        ):
            raise PipelineError(f"Dictionary {p.name}: object строка -> строка")
        extras.append(
            {"path": str(p.relative_to(ROOT)), "sha256": file_hash(p), "terms": data}
        )
    c["RESOLVED_DICTIONARIES"] = extras
    return c


def inside_root(value: str | Path) -> Path:
    p = Path(value)
    p = (ROOT / p).resolve() if not p.is_absolute() else p.resolve()
    if not p.is_relative_to(ROOT):
        raise PipelineError(f"Путь записи/входа должен находиться внутри {ROOT}")
    return p


def atomic_write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with NamedTemporaryFile(
            "w", encoding="utf-8", delete=False, dir=path.parent, suffix=".tmp"
        ) as handle:
            temporary = Path(handle.name)
            json.dump(payload, handle, ensure_ascii=False, indent=2, allow_nan=False)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def atomic_write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(
        "w", encoding="utf-8", delete=False, dir=path.parent, suffix=".tmp"
    ) as h:
        temp = Path(h.name)
        h.write(value)
        h.flush()
        os.fsync(h.fileno())
    try:
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def source_identity(path: Path) -> dict:
    before = path.stat()
    sha = file_hash(path)
    after = path.stat()
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        raise PipelineError("Source изменился во время hashing")
    return {
        "path": str(path.resolve()),
        "size": after.st_size,
        "mtime_ns": after.st_mtime_ns,
        "sha256": sha,
    }


def model_path(repo: str, override: str | None, required: tuple[str, ...]) -> Path:
    if override:
        candidates = [Path(override).expanduser().resolve()]
    else:
        hf_home = Path(
            os.environ.get("HF_HOME", Path.home() / ".cache" / "huggingface")
        )
        hub = Path(os.environ.get("HF_HUB_CACHE", hf_home / "hub"))
        parent = hub / ("models--" + repo.replace("/", "--")) / "snapshots"
        candidates = sorted(
            parent.glob("*"), key=lambda p: p.stat().st_mtime_ns, reverse=True
        )
    for p in candidates:
        if all((p / filename).is_file() for filename in required):
            return p
    raise PipelineError(
        f"Модель {repo} отсутствует/неполна в локальном cache. "
        "Проверьте HF доступ и локальные файлы; автоматическая загрузка отключена."
    )


def model_metadata(path: Path) -> dict:
    files = {}
    for p in sorted(path.iterdir()):
        if p.is_file():
            stat = p.stat()
            files[p.name] = {"size": stat.st_size, "mtime_ns": stat.st_mtime_ns}
            if stat.st_size < 5_000_000:
                files[p.name]["sha256"] = file_hash(p)
    return {"path": str(path), "revision": path.name, "files": files}


def runtime_info(config: dict) -> dict:
    import torch

    torch.set_num_threads(config["CPU_THREADS"])
    cuda = torch.cuda.is_available()
    if config["DEVICE"] == "cuda" and not cuda:
        raise PipelineError(
            "CUDA запрошена, но недоступна; используйте --device cpu явно"
        )
    device = "cuda" if cuda and config["DEVICE"] != "cpu" else "cpu"
    return {
        "device": device,
        "compute_type": ("float16" if config["GPU_FP16"] else "float32")
        if device == "cuda"
        else "int8",
        "gpu": torch.cuda.get_device_name(0) if cuda else None,
        "cuda": torch.version.cuda,
        "torch": torch.__version__,
        "translation_device": config["TRANSLATION_DEVICE"],
    }


def processing_metadata(config: dict, runtime: dict, paths: dict) -> dict:
    excluded = {"INPUT", "OUTPUT", "TEMP", "SPEAKER_ROLES", "JSON_SAVE_INTERVAL"}
    effective = {k: v for k, v in config.items() if k not in excluded}
    versions = {}
    for pkg in (
        "faster-whisper",
        "ctranslate2",
        "torch",
        "pyannote.audio",
        "transformers",
        "tokenizers",
        "librosa",
    ):
        versions[pkg] = importlib.metadata.version(pkg)
    ffmpeg = subprocess.run(
        ["ffmpeg", "-version"], capture_output=True, text=True, check=True
    ).stdout.splitlines()[0]
    return {
        "pipeline_version": VERSION,
        "schema_version": SCHEMA,
        "stage_policy": __import__("translation_v156").stage_policy(
            sys.modules[__name__]
        ),
        "algorithm_sha256": file_hash(Path(__file__)),
        "config": effective,
        "models": {name: model_metadata(p) for name, p in paths.items()},
        "libraries": versions,
        "runtime": runtime,
        "audio": {
            "filter_chain": "loudnorm,highpass=f=80,lowpass=f=8000",
            "rate": 16000,
            "channels": 1,
            "ffmpeg": ffmpeg,
        },
        "language_policy": 3,
        "mapping_version": 3,
        "translation_channels": ["ru", "en"],
        "translation_group_policy": 1,
        "build": BUILD,
        "text_lid": {
            "version": "langid-1.1.6",
            "sha256": file_hash(ROOT / "vendor/text_lid/langid/langid.py"),
        },
    }


def validate_interval(item: dict, duration: float) -> None:
    a, b = item.get("start"), item.get("end")
    if any(
        isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v)
        for v in (a, b)
    ):
        raise ValueError("invalid interval type")
    if not 0 <= a < b <= duration + 0.001:
        raise ValueError("invalid interval range")


def validate_chunks(text: str, chunks: list[dict], target: str = "ru") -> None:
    position = 0
    for ch in chunks:
        if type(ch.get("start")) is not int or type(ch.get("end")) is not int:
            raise ValueError("invalid source span")
        if ch["start"] != position or not position < ch["end"] <= len(text):
            raise ValueError("translation coverage gap/overlap")
        if ch.get("source_text") != text[position : ch["end"]]:
            raise ValueError("translation source modified")
        if ch.get("status") not in TRANSLATION_STATUSES:
            raise ValueError("invalid chunk status")
        if ch.get("translation_" + target) is not None and not isinstance(
            ch["translation_" + target], str
        ):
            raise ValueError("invalid chunk translation")
        if (
            type(ch.get("source_token_count")) is not int
            or ch["source_token_count"] < 0
        ):
            raise ValueError("invalid token count")
        if ch.get("children"):
            validate_chunks(ch["source_text"], ch["children"], target)
        position = ch["end"]
    if position != len(text):
        raise ValueError("translation source coverage incomplete")


def validate_cache(data: dict, source: dict, fingerprint: str) -> None:
    if (
        not isinstance(data, dict)
        or data.get("schema_version") != SCHEMA
        or data.get("pipeline_version") != VERSION
    ):
        raise ValueError("incompatible schema/version")
    if (
        data.get("source") != source
        or data.get("processing_fingerprint") != fingerprint
    ):
        raise ValueError("incompatible source/processing")
    if digest(data.get("processing")) != fingerprint:
        raise ValueError("processing metadata hash mismatch")
    if data.get("state_sha256") != digest(
        {k: v for k, v in data.items() if k != "state_sha256"}
    ):
        raise ValueError("cache payload checksum mismatch")
    duration = data.get("duration")
    if (
        isinstance(duration, bool)
        or not isinstance(duration, (float, int))
        or not math.isfinite(duration)
        or duration <= 0
    ):
        raise ValueError("invalid duration")
    for key in ("diarization", "units", "segments"):
        if not isinstance(data.get(key), list):
            raise ValueError(f"invalid {key}")
    if type(data.get("diarization_done")) is not bool:
        raise ValueError("invalid diarization stage")
    for item in data["diarization"]:
        validate_interval(item, duration)
        if not isinstance(item.get("speaker_id"), str):
            raise ValueError("invalid speaker")
    ids = set()
    for seg in data["segments"]:
        validate_interval(seg, duration)
        if not isinstance(seg.get("segment_id"), str) or seg["segment_id"] in ids:
            raise ValueError("duplicate/invalid segment id")
        ids.add(seg["segment_id"])
        if not isinstance(seg.get("raw_text"), str) or seg.get(
            "raw_sha256"
        ) != text_hash(seg["raw_text"]):
            raise ValueError("raw text integrity violation")
        if seg.get("translation_status") not in TRANSLATION_STATUSES:
            raise ValueError("invalid translation status")
        for field in ("speaker_id", "language", "language_method", "corrected_text"):
            if not isinstance(seg.get(field), str):
                raise ValueError(f"invalid segment {field}")
        if not isinstance(seg.get("chunks"), list) or not isinstance(
            seg.get("warnings"), list
        ):
            raise ValueError("invalid segment collections")
        if (
            seg["translation_status"] == "complete"
            and seg.get("semantic_validation") != "reviewed"
        ):
            raise ValueError("complete translation lacks semantic review")
        if seg["language"] == "ru" and seg.get("translation_ru"):
            raise ValueError("Russian source must not be translated")
        if seg.get("transcription_status") not in {
            "complete",
            "needs_review",
            "empty",
            "failed",
            "unsupported_language",
        }:
            raise ValueError("invalid transcription status")
        if seg.get("chunks"):
            validate_chunks(seg["raw_text"], seg["chunks"])
    unit_ids = set()
    previous_end = 0.0
    referenced = []
    for unit in data["units"]:
        validate_interval(unit, duration)
        if unit.get("unit_id") in unit_ids or type(unit.get("asr_done")) is not bool:
            raise ValueError("invalid ASR unit")
        unit_ids.add(unit["unit_id"])
        if abs(unit["start"] - previous_end) > 0.001:
            raise ValueError("ASR timeline gap/overlap")
        previous_end = unit["end"]
        if not isinstance(unit.get("segment_ids"), list) or not set(
            unit["segment_ids"]
        ).issubset(ids):
            raise ValueError("unit points to missing ASR")
        if unit["asr_done"] and not unit["segment_ids"]:
            raise ValueError("completed unit without result")
        referenced.extend(unit["segment_ids"])
        if unit["asr_done"] and unit.get("raw_pieces_sha256") != digest(
            unit.get("raw_pieces")
        ):
            raise ValueError("raw ASR pieces integrity violation")
    if len(referenced) != len(set(referenced)) or set(referenced) != ids:
        raise ValueError("orphan/duplicated ASR results")
    if data["units"] and abs(previous_end - duration) > 0.001:
        raise ValueError("ASR timeline tail missing")
    validate_translation_groups(data)


def cache_location(base: Path, source: dict, processing: dict) -> tuple[Path, str]:
    fingerprint = digest(processing)
    folder = base / (source["sha256"][:20] + "-" + fingerprint[:20])
    return folder, fingerprint


def migrate_ux_predecessor(base: Path, source: dict, processing: dict) -> dict | None:
    import translation_v156 as mt

    return mt.reuse_stages(sys.modules[__name__], base, source, processing)


def read_cache(
    folder: Path, source: dict, fingerprint: str
) -> tuple[dict | None, Path]:
    candidates = [
        folder,
        *sorted(
            folder.parent.glob(folder.name + "-recovery-*"),
            key=lambda f: f.stat().st_mtime_ns,
            reverse=True,
        ),
    ]
    corrupt = False
    for candidate in candidates:
        path = candidate / "resume.json"
        if not path.exists():
            continue
        try:
            data = read_json(path)
            validate_cache(data, source, fingerprint)
            return data, candidate
        except (PipelineError, ValueError, KeyError, TypeError, AttributeError):
            corrupt = True
            LOG.warning(
                "Cache miss: incompatible/corrupt state; original retained: %s", path
            )
    return None, folder.with_name(
        folder.name + "-recovery-" + uuid.uuid4().hex[:8]
    ) if corrupt else folder


@contextmanager
def run_lock(folder: Path):
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / "run.lock"
    try:
        handle = path.open("x", encoding="utf-8")
    except FileExistsError as error:
        raise PipelineError(
            f"Cache занят или прежний запуск аварийно завершился: {path}. "
            "Проверьте PID в lock; используйте --fresh для нового namespace."
        ) from error
    try:
        handle.write(json.dumps({"pid": os.getpid(), "created_at": utc_now()}))
        handle.close()
        yield
    finally:
        handle.close()
        path.unlink(missing_ok=True)


def prepare_audio(source: Path, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name("prepared-" + uuid.uuid4().hex + ".wav")
    command = [
        "ffmpeg",
        "-nostdin",
        "-v",
        "error",
        "-y",
        "-i",
        str(source),
        "-af",
        "loudnorm,highpass=f=80,lowpass=f=8000",
        "-ar",
        "16000",
        "-ac",
        "1",
        str(tmp),
    ]
    try:
        result = subprocess.run(command, capture_output=True, text=True, check=False)
        if result.returncode or not tmp.is_file() or tmp.stat().st_size < 44:
            raise PipelineError(
                "FFmpeg не смог подготовить WAV; проверьте формат и доступ к файлу"
            )
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def load_audio(path: Path):
    import librosa

    audio, sr = librosa.load(path, sr=16000, mono=True)
    if sr != 16000 or not len(audio):
        raise PipelineError("Пустой/неверный WAV")
    return audio


def normalize_diarization(
    output: Any, duration: float
) -> tuple[list[dict], list[dict]]:
    regular = getattr(output, "speaker_diarization", output)
    exclusive = getattr(output, "exclusive_speaker_diarization", None)

    def extract(annotation):
        if not hasattr(annotation, "itertracks"):
            raise PipelineError("pyannote вернула неподдерживаемый output")
        result = []
        for turn, _, speaker in annotation.itertracks(yield_label=True):
            a, b = max(0.0, float(turn.start)), min(duration, float(turn.end))
            if not math.isfinite(a) or not math.isfinite(b):
                raise PipelineError("pyannote interval is non-finite")
            if b > a:
                result.append({"start": a, "end": b, "speaker_id": str(speaker)})
        return sorted(result, key=lambda x: (x["start"], x["end"], x["speaker_id"]))

    return extract(regular), extract(exclusive) if exclusive is not None else []


def diarize(audio: Any, config: dict, path: Path) -> tuple[list[dict], list[dict]]:
    import torch
    from pyannote.audio import Pipeline

    try:
        pipeline = Pipeline.from_pretrained(str(path / "config.yaml"))
        if pipeline is None:
            raise PipelineError(
                "pyannote model unavailable: проверьте cached gated dependencies"
            )
        options = {"min_speakers": config["MIN_SPEAKERS"]}
        if config["MAX_SPEAKERS"] is not None:
            options["max_speakers"] = config["MAX_SPEAKERS"]
        output = pipeline(
            {"waveform": torch.from_numpy(audio).unsqueeze(0), "sample_rate": 16000},
            **options,
        )
        return normalize_diarization(output, len(audio) / 16000)
    except PipelineError:
        raise
    except Exception as error:
        raise PipelineError(
            f"pyannote: {type(error).__name__}; проверьте локальные модели/доступ HF"
        ) from error


def speaker_for(
    start: float, end: float, intervals: list[dict]
) -> tuple[str, list[str]]:
    scores = {}
    for item in intervals:
        overlap = max(0.0, min(end, item["end"]) - max(start, item["start"]))
        if overlap:
            scores[item["speaker_id"]] = scores.get(item["speaker_id"], 0.0) + overlap
    if not scores:
        return "UNKNOWN", []
    ordered = sorted(scores, key=lambda s: (-scores[s], s))
    return ordered[0], ordered if len(ordered) > 1 else []


def initial_units(audio: Any, config: dict) -> list[dict]:
    """Partition the FULL timeline; pyannote cannot silently drop an uncovered voice."""
    import numpy as np

    duration = len(audio) / 16000
    result = []
    start = 0.0
    while start < duration - 1e-6:
        end = min(duration, start + config["ASR_UNIT_SECONDS"])
        if end < duration:
            # Choose a low-energy 50ms point before the cap; no timeline gaps.
            candidates = np.arange(max(start + 1, end - 1.5), end, 0.05)
            if len(candidates):
                end = float(
                    min(
                        candidates,
                        key=lambda t: float(
                            np.mean(
                                audio[int(t * 16000) : int((t + 0.05) * 16000)] ** 2
                            )
                        ),
                    )
                )
        result.append({"start": start, "end": end})
        start = end
    return result


class ASRBackend(Protocol):
    """Optional backend seam. External Zulu implementations must return the same contract."""

    def detect(self, audio: Any) -> tuple[str, float]: ...
    def transcribe(self, audio: Any, language: str | None) -> tuple[list[Any], Any]: ...


class UnsupportedZuluBackend:
    def detect(self, audio: Any):
        raise PipelineError("Zulu ASR unsupported by current backend")

    def transcribe(self, audio: Any, language: str | None):
        raise PipelineError("Zulu ASR unsupported by current backend")


class WhisperBackend:
    def __init__(self, path: Path, config: dict, runtime: dict):
        from faster_whisper import WhisperModel

        self.config = config
        try:
            self.model = WhisperModel(
                str(path),
                device=runtime["device"],
                compute_type=runtime["compute_type"],
                local_files_only=True,
            )
        except Exception as error:
            raise PipelineError(
                f"Whisper {runtime['device']}: {type(error).__name__}. "
                "Проверьте CUDA/VRAM; CPU запуск: --device cpu"
            ) from error

    def detect(self, audio: Any) -> tuple[str, float]:
        language, probability, _ = self.model.detect_language(
            audio=audio, vad_filter=False
        )
        return language, float(probability)

    def transcribe(self, audio: Any, language: str | None):
        c = self.config
        pieces, info = self.model.transcribe(
            audio,
            language=language,
            task="transcribe",
            beam_size=c["BEAM_SIZE"],
            best_of=c["BEST_OF"],
            temperature=c["TEMPERATURE"],
            vad_filter=c["USE_VAD"],
            vad_parameters={
                "threshold": c["VAD_THRESHOLD"],
                "min_speech_duration_ms": c["VAD_MIN_SPEECH_MS"],
                "min_silence_duration_ms": c["VAD_MIN_SILENCE_MS"],
                "speech_pad_ms": c["VAD_PAD_MS"],
            },
            condition_on_previous_text=c["CONDITION_ON_PREVIOUS_TEXT"],
            repetition_penalty=c["REPETITION_PENALTY"],
            no_repeat_ngram_size=c["NO_REPEAT_NGRAM_SIZE"],
            compression_ratio_threshold=c["COMPRESSION_RATIO_THRESHOLD"],
            log_prob_threshold=c["LOG_PROB_THRESHOLD"],
            no_speech_threshold=c["NO_SPEECH_THRESHOLD"],
            word_timestamps=True,
        )
        return list(pieces), info


def language_units(audio: Any, backend: ASRBackend, config: dict) -> list[dict]:
    units = []

    def visit(start: float, end: float):
        force = config["FORCE_LANGUAGE"]
        probes = []
        if force:
            language, confidence, method, candidate = force, None, "forced", force
        else:
            probe = config["LANGUAGE_PROBE_SECONDS"]
            cursor = start
            while cursor < end - 1e-6:
                stop = min(end, cursor + probe)
                detected = backend.detect(
                    audio[int(cursor * 16000) : int(stop * 16000)]
                )
                probes.append(
                    {
                        "start": cursor,
                        "end": stop,
                        "language": detected[0],
                        "confidence": detected[1],
                    }
                )
                cursor = stop
            threshold = config["LANGUAGE_CONFIDENCE"]
            strong = [item for item in probes if item["confidence"] >= threshold]
            strong_codes = {item["language"] for item in strong}
            minimum = config["MIN_LANGUAGE_UNIT_SECONDS"]
            for left, right in zip(strong, strong[1:]):
                if left["language"] != right["language"] and end - start >= 2 * minimum:
                    cut = min(
                        end - minimum,
                        max(start + minimum, (left["end"] + right["start"]) / 2),
                    )
                    visit(start, cut)
                    visit(cut, end)
                    return
            best = max(probes, key=lambda item: item["confidence"])
            candidate, confidence = best["language"], best["confidence"]
            language = (
                candidate
                if confidence >= threshold and len(strong_codes) == 1
                else "unknown"
            )
            method = "whisper_audio_local"
            if language == "unknown":
                # Use more audio for detection only, never inherit a speaker's language.
                pad = 3.0 if end - start < config["MIN_LANGUAGE_UNIT_SECONDS"] else 0.0
                ca, cb = max(0, start - pad), min(len(audio) / 16000, end + pad)
                context = backend.detect(audio[int(ca * 16000) : int(cb * 16000)])
                probes.append(
                    {
                        "start": ca,
                        "end": cb,
                        "language": context[0],
                        "confidence": context[1],
                        "method": "context",
                    }
                )
                if (
                    context[1] >= threshold
                    and len(strong_codes) < 2
                    and (context[0] == candidate or not strong)
                ):
                    candidate, confidence = context
                    language = candidate
                    method = (
                        "whisper_audio_expanded_context"
                        if pad
                        else "whisper_audio_full_unit"
                    )
        units.append(
            {
                "unit_id": f"u{len(units):05}",
                "start": start,
                "end": end,
                "language": language,
                "language_confidence": confidence,
                "language_method": method,
                "candidate_language": candidate,
                "probes": probes,
                "asr_done": False,
                "segment_ids": [],
                "raw_pieces": [],
            }
        )

    for span in initial_units(audio, config):
        visit(span["start"], span["end"])
    return units


class TextLanguageValidator:
    """Local byte-ngram classifier; scores are evidence, not calibrated accuracy."""

    def __init__(self):
        location = str(ROOT / "vendor/text_lid")
        if location not in sys.path:
            sys.path.insert(0, location)
        from langid.langid import LanguageIdentifier, model

        self.model = LanguageIdentifier.from_modelstring(model, norm_probs=True)

    def assess(self, text: str) -> dict:
        ranked = self.model.rank(text)
        top = [(code, float(score)) for code, score in ranked[:3]]
        enough = (
            len(re.findall(r"\w+", text)) >= 5 and sum(c.isalpha() for c in text) >= 25
        )
        return {
            "model": "langid-1.1.6-full97",
            "language": top[0][0],
            "score": top[0][1],
            "margin": top[0][1] - top[1][1],
            "top": top,
            "strong": enough and top[0][1] >= 0.95 and top[0][1] - top[1][1] >= 0.5,
        }


def resolve_text_language(
    detected: str, evidence: dict, continuity: bool, forced: bool = False
) -> tuple[str, str]:
    if forced:
        return detected, "explicit_force"
    text_language = evidence["language"]
    # Only reconcile the demonstrated Latin-language confusions. Other scripts
    # retain the existing conservative policy; unsupported text is not English.
    if (
        detected in {"af", "nl", "en"}
        and evidence["strong"]
        and text_language in {"af", "nl", "en"}
    ):
        if detected != text_language:
            reason = "statistical_text_conflict"
            if detected == "nl" and text_language == "af" and continuity:
                reason = "afrikaans_text_and_continuity"
            return text_language, reason
    return (
        detected,
        "audio_text_agree"
        if detected == text_language
        else "insufficient_text_evidence",
    )


def reconcile_unit(
    unit: dict,
    segments: list[dict],
    units: list[dict],
    index: int,
    audio: Any,
    backend: ASRBackend,
    lid: Any,
    config: dict,
) -> None:
    text = "".join(s["raw_text"] for s in segments)
    evidence = lid.assess(text)
    continuity = (
        0 < index < len(units) - 1
        and units[index - 1]["language"] == units[index + 1]["language"] == "af"
        and unit["end"] - unit["start"] <= 12
    )
    resolved, reason = resolve_text_language(
        unit["language"], evidence, continuity, bool(config["FORCE_LANGUAGE"])
    )
    provenance = {
        "detected_language": unit["language"],
        "text_evidence": evidence,
        "af_neighbors": continuity,
        "resolved_language": resolved,
        "reason": reason,
        "raw_selection": "primary_preserved",
    }
    # Exactly one alternative on a demonstrated conflict or AF/NL continuity.
    # Never replace immutable primary ASR by an alternative merely for fluency.
    alternative_language = (
        resolved
        if resolved != unit["language"]
        else "af"
        if unit["language"] == "nl" and continuity
        else None
    )
    if alternative_language and not config["FORCE_LANGUAGE"]:
        a = max(0, unit["start"] - config["BOUNDARY_PADDING_SECONDS"])
        b = min(len(audio) / 16000, unit["end"] + config["BOUNDARY_PADDING_SECONDS"])
        try:
            pieces, info = backend.transcribe(
                audio[int(a * 16000) : int(b * 16000)], alternative_language
            )
            alt_text = "".join(piece.text for piece in pieces)
            provenance["alternative"] = {
                "requested_language": alternative_language,
                "text_evidence": lid.assess(alt_text),
                "pieces": [
                    {
                        "start": a + x.start,
                        "end": a + x.end,
                        "text": x.text,
                        "tokens": list(x.tokens),
                        "avg_logprob": float(x.avg_logprob),
                    }
                    for x in pieces
                ],
                "info_language": info.language,
            }
        except Exception as error:
            provenance["alternative"] = {
                "requested_language": alternative_language,
                "status": "failed",
                "error_type": type(error).__name__,
            }
            for seg in segments:
                seg["warnings"].append("alternative_language_check_failed")
                seg["transcription_status"] = "needs_review"
    unit["language_reconciliation"] = provenance
    for seg in segments:
        local = lid.assess(seg["raw_text"])
        # Small fragments use their own audio unit's text, never the speaker's language.
        local_safe = local["strong"] and (
            local["language"] == evidence["language"]
            or len(seg["raw_text"].split()) >= 8
            and {local["language"], evidence["language"]} != {"af", "nl"}
        )
        selected = local if local_safe else evidence
        language, why = resolve_text_language(
            seg["language"], selected, continuity, bool(config["FORCE_LANGUAGE"])
        )
        seg.update(
            detected_language=seg["language"],
            detected_language_confidence=seg["language_confidence"],
            resolved_language_score=selected["score"],
            resolved_language=language,
            language=language,
            language_resolution={
                "reason": why,
                "evidence": selected,
                "scope": "segment" if local_safe else "asr_unit",
            },
        )
        seg["corrected_text"] = corrected(seg["raw_text"], language, config)
        if language != seg["detected_language"]:
            seg.setdefault("diagnostics", []).append("audio_text_language_reconciled")
        if (
            why == "insufficient_text_evidence"
            and selected["score"] >= 0.75
            and selected["language"] != seg["language"]
            and len(text.split()) >= 5
        ):
            seg["warnings"].append("uncertain_language_conflict")
            seg["transcription_status"] = "needs_review"
        if continuity and seg["detected_language"] == "nl":
            seg.setdefault("diagnostics", []).append(
                "af_nl_continuity_primary_text_retained"
            )


def corrected(text: str, language: str, config: dict) -> str:
    # Only same-script substitutions; no Chinese -> Russian replacement before MT.
    terms = {**config["TCM_DICT"], **config["TERMS"]}
    for item in config.get("RESOLVED_DICTIONARIES", []):
        terms.update(item["terms"])
    applicable = {}
    for old, new in terms.items():
        if (
            language == "zh"
            and re.search(r"[\u4e00-\u9fff]", old)
            and not re.search(r"[А-Яа-яЁё]", new)
        ):
            applicable[old] = new
        elif (
            language == "ru"
            and re.search(r"[А-Яа-яЁё]", old)
            and not re.search(r"[\u4e00-\u9fff]", old)
        ):
            applicable[old] = new
    if not applicable:
        return text
    pattern = "|".join(re.escape(k) for k in sorted(applicable, key=len, reverse=True))
    return re.sub(pattern, lambda m: applicable[m.group(0)], text)


def script_language(text: str, detected: str, forced: bool) -> tuple[str, list[str]]:
    if forced:
        return detected, []
    counts = {
        "ru": len(re.findall(r"[А-Яа-яЁё]", text)),
        "he": len(re.findall(r"[\u0590-\u05ff]", text)),
        "zh": len(re.findall(r"[\u4e00-\u9fff]", text)),
        "latin": len(re.findall(r"[A-Za-zÀ-ž]", text)),
    }
    present = [k for k, v in counts.items() if v >= 2]
    if len(present) > 1:
        return "mixed", ["mixed_script_review"]
    if present and present[0] != "latin" and present[0] != detected:
        # Cyrillic alone cannot prove Russian (nor Hebrew script prove Hebrew).
        return "unknown", ["audio_text_language_disagreement"]
    return detected, []


def make_segment(
    unit: dict,
    index: int,
    start: float,
    end: float,
    text: str,
    speaker: str,
    overlap: list[str],
    config: dict,
    words: list | None = None,
    warning: str | None = None,
) -> dict:
    language, warnings = script_language(
        text, unit["language"], unit["language_method"] == "forced"
    )
    if overlap:
        warnings.append("overlapping_speakers")
    if warning:
        warnings.append(warning)
    if speaker == "UNKNOWN":
        warnings.append("speaker_unassigned")
    status = "complete"
    if not text.strip():
        status = "empty"
        warnings.append("empty_asr_audio_retained")
    elif warnings or language in {"unknown", "mixed"}:
        status = "needs_review"
    return {
        "segment_id": f"{unit['unit_id']}-s{index:04}",
        "asr_unit_id": unit["unit_id"],
        "start": start,
        "end": end,
        "speaker_id": speaker,
        "speaker_role": None,
        "language": language,
        "language_confidence": unit["language_confidence"],
        "language_method": unit["language_method"],
        "raw_text": text,
        "raw_sha256": text_hash(text),
        "corrected_text": corrected(text, language, config),
        "translation_ru": None,
        "translation_corrected_ru": None,
        "transcription_status": status,
        "translation_status": "pending",
        "semantic_validation": "unreviewed",
        "warnings": warnings,
        "overlap_speakers": overlap,
        "words": words or [],
        "chunks": [],
    }


def transcribe_unit(
    unit: dict,
    audio: Any,
    backend: ASRBackend,
    config: dict,
    intervals: list[dict],
    exclusive: list[dict] | None = None,
) -> list[dict]:
    duration = len(audio) / 16000
    a = max(0.0, unit["start"] - config["BOUNDARY_PADDING_SECONDS"])
    b = min(duration, unit["end"] + config["BOUNDARY_PADDING_SECONDS"])
    candidate = unit["candidate_language"]
    # A low-confidence candidate is explicit in metadata, never a hidden English fallback.
    pieces, info = backend.transcribe(audio[int(a * 16000) : int(b * 16000)], candidate)
    alternative = []
    compare_warning = False
    if config["VAD_COMPARE"] and isinstance(backend, WhisperBackend):
        old_vad = backend.config["USE_VAD"]
        try:
            backend.config["USE_VAD"] = not old_vad
            other, other_info = backend.transcribe(
                audio[int(a * 16000) : int(b * 16000)], candidate
            )
            alternative = [
                {
                    "start": a + x.start,
                    "end": a + x.end,
                    "text": x.text,
                    "tokens": list(x.tokens),
                }
                for x in other
            ]

            def normalized(text):
                return re.sub(r"\W+", "", text.lower())

            compare_warning = normalized(
                " ".join(x.text for x in pieces)
            ) != normalized(" ".join(x.text for x in other))
            unit["vad_comparison"] = {
                "primary_vad": old_vad,
                "alternative_vad": not old_vad,
                "different": compare_warning,
                "alternative_pieces": alternative,
                "alternative_duration_after_vad": float(other_info.duration_after_vad),
            }
        finally:
            backend.config["USE_VAD"] = old_vad
    unit["asr_info"] = {
        "language": info.language,
        "language_probability": float(info.language_probability),
        "duration_after_vad": float(info.duration_after_vad),
    }
    segments = []
    unit["raw_pieces"] = []
    for piece in pieces:
        words = [
            {
                "start": max(0.0, a + w.start),
                "end": min(duration, a + w.end),
                "word": w.word,
                "probability": float(w.probability),
            }
            for w in (piece.words or [])
        ]
        unit["raw_pieces"].append(
            {
                "start": a + piece.start,
                "end": a + piece.end,
                "text": piece.text,
                "tokens": list(piece.tokens),
                "words": words,
                "avg_logprob": float(piece.avg_logprob),
                "no_speech_prob": float(piece.no_speech_prob),
            }
        )
        owned = [
            w
            for w in words
            if unit["start"] <= (w["start"] + w["end"]) / 2 < unit["end"]
            or unit["end"] == duration
            and (w["start"] + w["end"]) / 2 == duration
        ]
        if words:
            groups = []
            for word in owned:
                speaker, _ = speaker_for(
                    word["start"],
                    max(word["end"], word["start"] + 0.001),
                    exclusive or intervals,
                )
                _, overlap = speaker_for(
                    word["start"], max(word["end"], word["start"] + 0.001), intervals
                )
                if groups and groups[-1][0] == speaker:
                    groups[-1][1].append(word)
                    groups[-1][2].update(overlap)
                else:
                    groups.append((speaker, [word], set(overlap)))
            for speaker, group, overlap in groups:
                start = min(unit["end"] - 0.001, max(unit["start"], group[0]["start"]))
                end = min(unit["end"], max(start + 0.001, group[-1]["end"]))
                if end <= start:
                    continue
                text = "".join(w["word"] for w in group)
                segments.append(
                    make_segment(
                        unit,
                        len(segments),
                        start,
                        end,
                        text,
                        speaker,
                        sorted(overlap),
                        config,
                        group,
                    )
                )
            if "".join(w["word"] for w in words).strip() != piece.text.strip():
                unit.setdefault("warnings", []).append(
                    "word_text_mismatch_original_in_raw_pieces"
                )
        elif piece.text.strip():
            # Missing alignment is never silently discarded. Keep raw piece with a warning.
            start, end = (
                max(unit["start"], a + piece.start),
                min(unit["end"], a + piece.end),
            )
            if end <= start:
                start, end = unit["start"], unit["end"]
            speaker, overlap = speaker_for(start, end, intervals)
            segments.append(
                make_segment(
                    unit,
                    len(segments),
                    start,
                    end,
                    piece.text,
                    speaker,
                    overlap,
                    config,
                    warning="missing_word_alignment_context_may_repeat",
                )
            )
    if not segments:
        speaker, overlap = speaker_for(unit["start"], unit["end"], intervals)
        segments.append(
            make_segment(
                unit, 0, unit["start"], unit["end"], "", speaker, overlap, config
            )
        )
    if compare_warning:
        for seg in segments:
            seg.setdefault("diagnostics", []).append("vad_alternative_differs")
    if unit.get("warnings"):
        for seg in segments:
            seg["transcription_status"] = (
                "needs_review" if seg["raw_text"].strip() else "empty"
            )
            seg["warnings"] += unit.get("warnings", [])
    return segments


def sentence_spans(text: str) -> list[tuple[int, int]]:
    """Keep every character, including whitespace; protect common abbreviations/decimals."""
    cuts = [0]
    for match in re.finditer(r"[.!?。！？]+(?:[\"'”’»)]*)(?:\s+|$)|[。！？]+", text):
        prefix = text[: match.start() + 1]
        if text[match.start()] == "." and (
            re.search(r"\b(?:Mr|Mrs|Ms|Dr|Prof|St|vs|etc)\.$", prefix, re.I)
            or re.search(r"\b[A-Z]\.$", prefix)
        ):
            continue
        cuts.append(match.end())
    cuts.append(len(text))
    cuts = sorted(set(cuts))
    return [(a, b) for a, b in zip(cuts, cuts[1:]) if b > a]


def split_budget(
    text: str, start: int, end: int, count: Callable[[str], int], budget: int
) -> list[tuple[int, int]]:
    if count(text[start:end]) <= budget:
        return [(start, end)]
    if end - start <= 1:
        raise PipelineError("Один символ превышает token budget")
    # Binary-search a character prefix, then prefer a clause/whitespace boundary.
    low, high = start + 1, end
    best = start
    while low <= high:
        mid = (low + high) // 2
        if count(text[start:mid]) <= budget:
            best, low = mid, mid + 1
        else:
            high = mid - 1
    if best <= start:
        raise PipelineError("Невозможно разместить source в token budget")
    boundaries = [
        m.end() + start for m in re.finditer(r"[,;:，；：]\s*|\s+", text[start:best])
    ]
    good = [p for p in boundaries if p - start >= (best - start) // 2]
    cut = good[-1] if good else best
    # Token counts need not be strictly monotonic; validate the actual chosen span.
    while count(text[start:cut]) > budget and cut > start + 1:
        cut -= 1
    return [(start, cut)] + split_budget(text, cut, end, count, budget)


def translation_spans(
    text: str, count: Callable[[str], int], budget: int
) -> list[tuple[int, int]]:
    result = []
    for start, end in sentence_spans(text):
        span = text[start:end]
        # Whisper often omits punctuation between thoughts. Split explicit clauses,
        # only for long sentences; do not merge them back into a single generation.
        cuts = [start]
        if count(span) > 32:
            cuts.extend(
                start + m.end()
                for m in re.finditer(
                    r"[,;:，；：]\s*|\s+(?=(?:So|so|But|but|And|and|Yes|yes)\s)", span
                )
            )
        cuts.append(end)
        for a, b in zip(sorted(set(cuts)), sorted(set(cuts))[1:]):
            if b > a:
                result.extend(split_budget(text, a, b, count, budget))
    return result


def semantic_translation_spans(
    text: str, count: Callable[[str], int], budget: int, language: str
) -> list[tuple[int, int]]:
    if language == "en":
        # Retain the measured English multi-thought omission repair.
        return translation_spans(text, count, budget)
    return [
        span
        for a, b in sentence_spans(text)
        for span in split_budget(text, a, b, count, budget)
    ]


def nllb_source(language: str, text: str, config: dict) -> tuple[str | None, list[str]]:
    if language == "zh":
        script = config["CHINESE_SCRIPT"]
        if script == "auto":
            # Conservative: shared Han characters do not prove Hans vs Hant.
            return "zho_Hans", ["chinese_script_assumed_Hans"]
        return "zho_" + script, []
    return NLLB.get(language), []


def load_nllb_tokenizer(path: Path):
    from transformers import NllbTokenizer

    # Transformers 5 AutoTokenizer may choose a generic backend whose src_lang
    # assignment does not update the encoder's language token.
    return NllbTokenizer.from_pretrained(
        str(path), local_files_only=True, src_lang="eng_Latn", legacy_behaviour=False
    )


class Translator:
    def __init__(self, path: Path, config: dict):
        import torch
        from transformers import AutoModelForSeq2SeqLM

        self.config = config
        self.target_language = "ru"
        self.device = config["TRANSLATION_DEVICE"]
        if self.device == "cuda" and not torch.cuda.is_available():
            raise PipelineError("NLLB CUDA недоступна; TRANSLATION_DEVICE=cpu")
        self.tokenizer = load_nllb_tokenizer(path)
        self.model = (
            AutoModelForSeq2SeqLM.from_pretrained(str(path), local_files_only=True)
            .to(self.device)
            .eval()
        )

    def set_language(self, source: str):
        token = self.tokenizer.convert_tokens_to_ids(source)
        if token is None or token == self.tokenizer.unk_token_id:
            raise PipelineError(f"NLLB language token unavailable: {source}")
        self.tokenizer.src_lang = source
        if self.tokenizer("", truncation=False)["input_ids"] != [
            token,
            self.tokenizer.eos_token_id,
        ]:
            raise PipelineError(
                "NLLB tokenizer did not apply the source language token"
            )

    def count(self, text: str) -> int:
        return len(self.tokenizer(text, truncation=False)["input_ids"])

    def generate(self, text: str) -> dict:
        import torch

        payload = self.tokenizer(text, return_tensors="pt", truncation=False)
        language_token = self.tokenizer.convert_tokens_to_ids(self.tokenizer.src_lang)
        if int(payload["input_ids"][0, 0]) != language_token:
            raise PipelineError("NLLB encoder language prefix mismatch")
        if payload["input_ids"].shape[-1] > self.config["TRANSLATION_TOKEN_BUDGET"]:
            raise PipelineError("NLLB input budget exceeded: no truncation permitted")
        payload = {k: v.to(self.device) for k, v in payload.items()}
        target_code = NLLB[self.target_language]
        target_token = self.tokenizer.convert_tokens_to_ids(target_code)
        if target_token is None or target_token == self.tokenizer.unk_token_id:
            raise PipelineError("NLLB target language token unavailable")
        with torch.inference_mode():
            result = self.model.generate(
                **payload,
                forced_bos_token_id=target_token,
                max_new_tokens=self.config["TRANSLATION_MAX_NEW_TOKENS"],
                num_beams=self.config["TRANSLATION_BEAMS"],
                do_sample=False,
                return_dict_in_generate=True,
            )
        ids = result.sequences[0].tolist()
        new_count = len(ids) - 1  # encoder-decoder starts with one decoder-start token.
        return {
            "text": self.tokenizer.decode(ids, skip_special_tokens=True),
            "source_language": self.tokenizer.src_lang,
            "target_language": self.target_language,
            "target_language_token_id": target_token,
            "input_language_token_id": language_token,
            "generated_ids": ids,
            "target_token_count": new_count,
            "eos": bool(ids and ids[-1] == self.tokenizer.eos_token_id),
            "hit_cap": new_count >= self.config["TRANSLATION_MAX_NEW_TOKENS"],
        }


def chunk_record(
    text: str, a: int, b: int, count: Callable[[str], int], target: str = "ru"
) -> dict:
    return {
        "start": a,
        "end": b,
        "source_text": text[a:b],
        "source_token_count": count(text[a:b]),
        "target_token_count": 0,
        "translation_" + target: None,
        "status": "pending",
        "warnings": [],
        "attempts": [],
        "generated_ids": [],
        "children": [],
    }


def pending_chunk(ch: dict) -> bool:
    return ch["status"] == "pending" or any(
        pending_chunk(child) for child in ch.get("children", [])
    )


def combine_children(ch: dict, target: str = "ru") -> None:
    children = ch["children"]
    field = "translation_" + target
    ch[field] = " ".join(x[field] or "" for x in children).strip() or None
    ch["status"] = (
        "needs_review"
        if all(
            x["status"] in {"needs_review", "complete", "not_required"}
            for x in children
        )
        else "partial"
    )
    ch["warnings"] = ["retried_smaller_chunks"] + sorted(
        {w for x in children for w in x["warnings"]}
    )


def degenerate_translation(source: str, target: str) -> bool:
    """Reject a demonstrated runaway loop, retaining the full attempt for review."""
    words = re.findall(r"\w+", target.casefold())
    source_words = re.findall(r"\w+", source.casefold())
    if len(words) < max(24, 3 * len(source_words)):
        return False
    for width in range(1, 9):
        for start in range(len(words) - 4 * width + 1):
            block = words[start : start + width]
            stop = start + width
            while words[stop : stop + width] == block:
                stop += width
            if stop - start >= max(16, 4 * width, len(words) // 2):
                return True
    return False


def generation_defects(
    source: str, result: dict, source_tokens: int, target_language: str = "ru"
) -> list[str]:
    target = result["text"]
    words = re.findall(r"[^\W\d_]+", source.casefold())
    latin = re.findall(r"[a-zÀ-ž]+", target.casefold())
    defects = []
    if re.search(r"([^\w\s])\1{5,}", target) or len(
        re.findall(r"[^\w\s]", target)
    ) > max(20, sum(c.isalpha() for c in target)):
        defects.append("pathological_punctuation")
    if degenerate_translation(source, target):
        defects.append("phrase_loop")
    ids = result.get("generated_ids", [])
    if len(ids) > max(32, 4 * source_tokens):
        for width in range(1, 9):
            if any(
                ids[i : i + width] * 8 == ids[i : i + width * 8]
                for i in range(max(0, len(ids) - width * 8 + 1))
            ):
                defects.append("token_loop")
                break
    if len(target) > max(160, 8 * len(source)) or result["target_token_count"] > max(
        96, 8 * source_tokens
    ):
        defects.append("abnormal_expansion")
    if len(words) >= 3 and sum(c.isalpha() for c in target) < 2:
        defects.append("near_empty_generation")
    # Preserve isolated names/acronyms. A copied multiword source run or a
    # predominantly Latin output is a failed RU generation, not an English fallback.
    copied = any(
        " ".join(words[i : i + 3]) in " ".join(latin) for i in range(len(words) - 2)
    )
    cyrillic = re.findall(r"[А-Яа-яЁё]+", target)
    if target_language == "ru" and (
        (copied and len(latin) >= 3) or (len(latin) >= 4 and len(latin) > len(cyrillic))
    ):
        defects.append("untranslated_source_or_mixed_generation")
    if target_language == "en" and len(cyrillic) >= 4 and len(cyrillic) > len(latin):
        defects.append("wrong_target_script")
    return defects


def translate_chunk(
    ch: dict,
    translator: Any,
    config: dict,
    depth: int = 0,
    checkpoint: Callable[[], None] = lambda: None,
    target_language: str = "ru",
) -> None:
    field = "translation_" + target_language
    if ch.get("children"):
        for child in ch["children"]:
            if pending_chunk(child):
                translate_chunk(
                    child, translator, config, depth + 1, checkpoint, target_language
                )
            combine_children(ch, target_language)
            checkpoint()
        return
    if not ch["source_text"].strip():
        ch.update(status="not_required")
        ch[field] = ""
        return
    try:
        result = translator.generate(ch["source_text"])
        warnings = []
        if not result["text"].strip():
            warnings.append("empty_generation")
        if result["hit_cap"]:
            warnings.append("output_cap")
        if not result["eos"]:
            warnings.append("missing_eos")
        defects = generation_defects(
            ch["source_text"], result, ch["source_token_count"], target_language
        )
        warnings.extend(defects)
        degenerate = bool(defects)
        if degenerate:
            warnings.append("degenerate_generation_preserved_in_attempts")
        target = result["target_token_count"]
        if ch["source_token_count"] >= 14 and target < 0.55 * ch["source_token_count"]:
            warnings.append("suspicious_shortening")
        src_numbers = re.findall(r"\d+", ch["source_text"])
        if any(number not in result["text"] for number in src_numbers):
            warnings.append("number_coverage_review")
        ch["attempts"].append({**result, "warnings": warnings, "depth": depth})
        ch.update(
            target_token_count=target,
            generated_ids=result["generated_ids"],
            warnings=warnings,
        )
        ch[field] = result["text"]
        ch["status"] = (
            "needs_review"  # never conflate EOS/coverage with semantic completeness
        )
        if not result["text"].strip():
            ch["status"] = "failed"
        if result["hit_cap"] or not result["eos"]:
            ch["status"] = "partial"
        if degenerate:
            # Raw source and the complete generated attempt remain available.
            ch.update(status="failed")
            ch[field] = None
        if (
            warnings
            and depth < config["TRANSLATION_RETRY_DEPTH"]
            and len(ch["source_text"].strip()) > 12
        ):
            text = ch["source_text"]
            midpoint = len(text) // 2
            boundaries = [m.end() for m in re.finditer(r"[,;:，；：]\s*|\s+", text)]
            cut = (
                min(boundaries, key=lambda x: abs(x - midpoint))
                if boundaries
                else midpoint
            )
            if 0 < cut < len(text):
                children = [
                    chunk_record(text, a, b, translator.count, target_language)
                    for a, b in [(0, cut), (cut, len(text))]
                ]
                ch["children"] = children
                ch["status"] = "partial"
                checkpoint()
                for child in children:
                    translate_chunk(
                        child,
                        translator,
                        config,
                        depth + 1,
                        checkpoint,
                        target_language,
                    )
                    combine_children(ch, target_language)
                    checkpoint()
    except Exception as error:
        ch["status"] = "failed"
        ch["warnings"] = ["translation_exception:" + type(error).__name__]


def build_translation_groups(segments: list[dict]) -> list[dict]:
    groups = []
    for seg in segments:
        previous = groups[-1] if groups else None
        join = previous is not None and (
            previous["language"] == seg["language"]
            and previous["speaker_id"] == seg["speaker_id"] != "UNKNOWN"
            and bool(seg["raw_text"].strip())
            and bool(previous["raw_text"].strip())
            and -0.05 <= seg["start"] - previous["end"] <= 1.25
            and seg["end"] - previous["start"] <= 24
            and not (
                re.search(r"[.!?。！？][\s\"']*$", previous["raw_text"])
                and len(previous["raw_text"].split()) >= 8
            )
        )
        if not join:
            previous = copy.deepcopy(seg)
            previous.update(
                group_id=f"t{len(groups):05}",
                member_ids=[],
                member_spans=[],
                raw_text="",
                chunks=[],
                translation_ru=None,
                translation_corrected_ru=None,
                translation_status="pending",
                warnings=[],
            )
            groups.append(previous)
        start = len(previous["raw_text"])
        previous["raw_text"] += seg["raw_text"]
        previous["member_ids"].append(seg["segment_id"])
        previous["member_spans"].append(
            {
                "segment_id": seg["segment_id"],
                "start": start,
                "end": len(previous["raw_text"]),
            }
        )
        previous["end"] = seg["end"]
        previous["raw_sha256"] = text_hash(previous["raw_text"])
        seg["translation_group_id"] = previous["group_id"]
    return groups


def sync_group_status(data: dict) -> None:
    groups = {g["group_id"]: g for g in data.get("translation_groups", [])}
    for seg in data["segments"]:
        group = groups.get(seg.get("translation_group_id"))
        if group:
            seg["translation_status"] = group["translation_status"]
            if group.get("semantic_validation") == "reviewed":
                seg["semantic_validation"] = "reviewed"
            else:
                seg["semantic_validation"] = "unreviewed"
            # Group-level MT has no claimed word/timestamp alignment.
            seg["translation_ru"] = None
            seg["translation_corrected_ru"] = None


def validate_translation_groups(data: dict) -> None:
    import translation_v156 as mt

    segments = {s["segment_id"]: s for s in data["segments"]}
    members = []
    group_ids = set()
    for group in data.get("translation_groups", []):
        gid = group["group_id"]
        if gid in group_ids:
            raise ValueError("duplicate translation group")
        group_ids.add(gid)
        ss = [segments[sid] for sid in group["member_ids"]]
        members.extend(group["member_ids"])
        if not ss or any(
            s["language"] != group["language"]
            or s["speaker_id"] != group["speaker_id"]
            or s.get("translation_group_id") != gid
            for s in ss
        ):
            raise ValueError("translation group crosses language/speaker")
        raw = "".join(s["raw_text"] for s in ss)
        if group["raw_text"] != raw or group["raw_sha256"] != text_hash(raw):
            raise ValueError("translation group source integrity")
        pos = 0
        for seg, span in zip(ss, group["member_spans"], strict=True):
            if span != {
                "segment_id": seg["segment_id"],
                "start": pos,
                "end": pos + len(seg["raw_text"]),
            }:
                raise ValueError("translation member span mismatch")
            pos += len(seg["raw_text"])
        if group["start"] != ss[0]["start"] or group["end"] != ss[-1]["end"]:
            raise ValueError("translation group timing mismatch")
        if group["translation_status"] not in TRANSLATION_STATUSES:
            raise ValueError("invalid group status")
        if (
            group["translation_status"] == "complete"
            and group.get("semantic_validation") != "reviewed"
        ):
            raise ValueError("unreviewed complete group")
        if group["chunks"]:
            validate_chunks(raw, group["chunks"])
        mt.validate_quality(sys.modules[__name__], group, "ru")
        en = group.get("english")
        if en is not None:
            if (
                en.get("raw_text") != raw
                or en.get("raw_sha256") != text_hash(raw)
                or en.get("language") != group["language"]
                or en.get("target_language") != "en"
            ):
                raise ValueError("English channel source/target mismatch")
            if en.get("translation_status") not in TRANSLATION_STATUSES:
                raise ValueError("invalid English channel status")
            if (
                en["translation_status"] == "complete"
                and en.get("semantic_validation") != "reviewed"
            ):
                raise ValueError("unreviewed complete English channel")
            if group["language"] == "en" and (
                en.get("translation_en") or en.get("chunks")
            ):
                raise ValueError("duplicate English translation")
            if en.get("translation_en") is not None and not isinstance(
                en["translation_en"], str
            ):
                raise ValueError("invalid English translation")
            if en.get("chunks"):
                validate_chunks(raw, en["chunks"], "en")
            mt.validate_quality(sys.modules[__name__], en, "en")
    if data.get("translation_groups") and members != list(segments):
        raise ValueError("translation groups must cover each segment once in order")


def translate_segment(
    seg: dict,
    translator: Any,
    config: dict,
    checkpoint: Callable[[], None] = lambda: None,
    retry: bool = False,
    target_language: str = "ru",
) -> None:
    if target_language not in {"ru", "en"}:
        raise PipelineError("Unsupported translation target")
    field = "translation_" + target_language
    raw = seg["raw_text"]
    if text_hash(raw) != seg["raw_sha256"]:
        raise PipelineError("Raw ASR integrity violation")
    if seg["language"] == target_language:
        seg.update(translation_status="not_required", chunks=[])
        seg[field] = None
        checkpoint()
        return
    if config["TRANSLATE_MODE"] == "off":
        seg.update(translation_status="disabled", chunks=[])
        seg[field] = None
        checkpoint()
        return
    source, warnings = nllb_source(seg["language"], raw, config)
    if not source:
        seg.update(translation_status="unsupported_language", chunks=[])
        seg[field] = None
        checkpoint()
        return
    if not raw.strip():
        seg.update(translation_status="failed", chunks=[])
        seg[field] = None
        checkpoint()
        return
    for warning in warnings:
        if warning not in seg["warnings"]:
            seg["warnings"].append(warning)
    translator.set_language(source)
    translator.target_language = target_language
    if not seg["chunks"]:
        seg["chunks"] = [
            chunk_record(raw, a, b, translator.count, target_language)
            for a, b in semantic_translation_spans(
                raw,
                translator.count,
                config["TRANSLATION_TOKEN_BUDGET"],
                seg["language"],
            )
        ]
    validate_chunks(raw, seg["chunks"], target_language)

    def save_progress():
        aggregate_translation(seg, target_language)
        checkpoint()

    def reset_failed(ch):
        if ch["status"] in {"failed", "partial"} and not ch.get("children"):
            ch["status"] = "pending"
        for child in ch.get("children", []):
            reset_failed(child)

    for ch in seg["chunks"]:
        if retry:
            reset_failed(ch)
        if pending_chunk(ch):
            translate_chunk(
                ch,
                translator,
                config,
                checkpoint=save_progress,
                target_language=target_language,
            )
            aggregate_translation(seg, target_language)
            checkpoint()
    aggregate_translation(seg, target_language)
    seg["translation_corrected_" + target_language] = corrected(
        seg[field] or "", target_language, config
    )
    checkpoint()


def aggregate_translation(seg: dict, target_language: str = "ru") -> None:
    chunks = seg["chunks"]
    field = "translation_" + target_language
    seg[field] = (
        " ".join(ch[field].strip() for ch in chunks if ch[field]).strip() or None
    )
    statuses = {ch["status"] for ch in chunks}
    if statuses <= {"needs_review", "complete", "not_required"}:
        seg["translation_status"] = "needs_review"
    elif seg[field]:
        seg["translation_status"] = "partial"
    else:
        seg["translation_status"] = "failed" if "failed" in statuses else "partial"
    seg["source_coverage"] = sum(ch["end"] - ch["start"] for ch in chunks) / max(
        1, len(seg["raw_text"])
    )


def format_time(seconds: float) -> str:
    total = max(0, int(seconds))
    return f"{total // 3600:02}:{total % 3600 // 60:02}:{total % 60:02}"


def segment_heading(seg: dict) -> str:
    role = (seg.get("speaker_role") + " | ") if seg.get("speaker_role") else ""
    return f"[{format_time(seg['start'])}–{format_time(seg['end'])}] {role}{seg['speaker_id']} | {seg['language'].upper()}"


def exported_status(data: dict) -> str:
    if not data["segments"] or any(not unit["asr_done"] for unit in data["units"]):
        return "partial"
    channels = [
        g["english"] for g in data.get("translation_groups", []) if "english" in g
    ]
    if any(
        c["translation_status"]
        in {"partial", "failed", "pending", "unsupported_language"}
        for c in channels
    ):
        return "partial"
    if any(
        seg["transcription_status"] in {"empty", "failed", "unsupported_language"}
        or seg["translation_status"]
        in {"partial", "failed", "pending", "unsupported_language"}
        for seg in data["segments"]
    ):
        return "partial"
    if any(c["translation_status"] == "needs_review" for c in channels) or any(
        seg["transcription_status"] == "needs_review"
        or seg["translation_status"] == "needs_review"
        for seg in data["segments"]
    ):
        return "needs_review"
    return "complete"


def group_needs_review(group: dict, segments: dict) -> bool:
    """Reading-copy warning for concrete issues, not the generic human-review gate."""
    if group.get("quality", {}).get("flags") or group.get("english", {}).get(
        "quality", {}
    ).get("flags"):
        return True
    if "semantic_review_required" in group.get("warnings", []):
        return True
    important = {
        "uncertain_language_conflict",
        "alternative_language_check_failed",
        "speaker_unassigned",
        "empty_asr_audio_retained",
        "mixed_script_review",
        "audio_text_language_disagreement",
        "missing_word_alignment_context_may_repeat",
    }
    members = [segments[sid] for sid in group.get("member_ids", [])] or [group]
    if any(
        s.get("transcription_status") in {"empty", "failed", "unsupported_language"}
        or s.get("speaker_id") == "UNKNOWN"
        or s.get("language") in {"unknown", "mixed"}
        or important.intersection(s.get("warnings", []))
        for s in members
    ):
        return True

    def chunk_issue(ch):
        if ch.get("children"):
            return any(chunk_issue(c) for c in ch["children"])
        return bool(ch.get("warnings")) or ch.get("status") in {
            "failed",
            "partial",
            "pending",
        }

    for target, record in (("ru", group), ("en", group.get("english", {}))):
        if group["language"] == target:
            continue
        if record.get("translation_status") in {
            "failed",
            "partial",
            "pending",
            "unsupported_language",
            "disabled",
        }:
            return True
        if not record.get("translation_" + target):
            return True
        if any(chunk_issue(c) for c in record.get("chunks", [])):
            return True
    return False


def export_all(data: dict, config: dict, out: Path) -> list[Path]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    out.mkdir(parents=True, exist_ok=True)
    payload = copy.deepcopy(data)
    payload.pop("state_sha256", None)
    payload["artifact_type"] = "transcript"
    payload["document_status"] = exported_status(data)
    payload["exported_at"] = utc_now()
    payload["semantic_note"] = (
        "Автоматический перевод не является подтверждённо полным по смыслу; needs_review до проверки человеком."
    )
    roles = config["SPEAKER_ROLES"].get(
        data["source"]["sha256"] + ":" + data["processing_fingerprint"], {}
    )
    for seg in payload["segments"]:
        seg["speaker_role"] = roles.get(seg["speaker_id"])
    for group in payload.get("translation_groups", []):
        group["speaker_role"] = roles.get(group["speaker_id"])
    # User output is flat. Cache identity never determines the publication path.
    stem = (
        Path(data["source"]["path"]).stem
        + ("_SMOKE" if data.get("processing", {}).get("limit_seconds") else "")
        + "_PRO"
    )
    paths = [out / (stem + ext) for ext in (".json", ".txt", ".docx")]
    payload["export_sha256"] = digest(payload)
    atomic_write_json(paths[0], payload)
    reading_groups = payload.get("translation_groups") or payload["segments"]
    segments_by_id = {s["segment_id"]: s for s in payload["segments"]}
    status = (
        "Требуется проверка отдельных фрагментов"
        if any(group_needs_review(g, segments_by_id) for g in reading_groups)
        or payload["document_status"] == "partial"
        else "Готово"
    )
    metadata = [
        f"Файл: {Path(data['source']['path']).name}",
        f"Длительность: {format_time(data['duration'])}",
        f"Языки: {', '.join(sorted({s['language'].upper() for s in data['segments']}))}",
        f"Статус: {status}",
    ]
    blocks = [f"Audio2Text CUDA PRO v{VERSION}", *metadata]
    doc = Document()
    doc.sections[0].top_margin = Inches(0.8)
    doc.sections[0].bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    for style_name in ("Title", "Normal"):
        style_properties = doc.styles[style_name].element.find(qn("w:pPr"))
        if style_properties is not None:
            for border in list(style_properties.findall(qn("w:pBdr"))):
                style_properties.remove(border)
    title = doc.add_paragraph(f"Audio2Text CUDA PRO v{VERSION}", style="Title")
    title.runs[0].font.size = Pt(20)
    title.runs[0].font.color.rgb = RGBColor.from_string("263445")
    for line in metadata:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(2)
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor.from_string("526071")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def paragraph(text: str, rtl: bool = False):
        p = doc.add_paragraph(text)
        if rtl:
            el = OxmlElement("w:bidi")
            el.set(qn("w:val"), "1")
            p._p.get_or_add_pPr().append(el)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p._p.get_or_add_pPr().find(qn("w:jc")).set(qn("w:val"), "start")
            for run in p.runs:
                r = OxmlElement("w:rtl")
                r.set(qn("w:val"), "1")
                run._r.get_or_add_rPr().append(r)
                fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
                fonts.set(qn("w:cs"), "Arial")
        return p

    for seg in reading_groups:
        heading = segment_heading(seg)
        blocks += ["", heading]
        group_paragraphs = []
        p = doc.add_paragraph()
        group_paragraphs.append(p)
        p.paragraph_format.space_before = Pt(7)
        speaker = seg["speaker_id"]
        if seg.get("speaker_role"):
            speaker += " · " + seg["speaker_role"]
        run = p.add_run(speaker)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("263445")
        r = p.add_run("   •   " + seg["language"].upper())
        r.bold = True
        badges = {
            "af": "218739",
            "nl": "D86B16",
            "en": "743EB0",
            "zh": "CF3038",
            "de": "B77C00",
            "fr": "9738A5",
            "it": "14866B",
            "he": "008B8B",
            "ru": "00A6D6",
        }
        r.font.color.rgb = RGBColor.from_string(badges.get(seg["language"], "526071"))
        r = p.add_run(f"   •   {format_time(seg['start'])}–{format_time(seg['end'])}")
        r.font.color.rgb = RGBColor.from_string("526071")
        r.font.size = Pt(9)

        def channel(label, text, color, rtl=False):
            blocks.extend([f"[{label}]", text])
            p = paragraph("", False)
            r = p.add_run(label.replace("/", " · "))
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string(color)
            p.paragraph_format.space_after = Pt(1)
            group_paragraphs.append(p)
            p = paragraph(text, rtl)
            for run in p.runs:
                run.font.color.rgb = RGBColor.from_string(color)
            group_paragraphs.append(p)

        channel(
            "ORIGINAL/" + seg["language"].upper(),
            seg["raw_text"].strip()
            or "[Нет распознанного текста; аудиоинтервал сохранён]",
            "1565C0",
            seg["language"] == "he",
        )
        if seg["language"] != "ru" and seg.get("translation_ru"):
            channel("RU", seg["translation_ru"], "00A6D6")
        en = seg.get("english", {})
        if seg["language"] != "en" and en.get("translation_en"):
            channel("EN", en["translation_en"], "D88A00")
        # Dictionary overlays remain available without replacing the original.
        original_corrected = "".join(
            segments_by_id[sid]["corrected_text"] for sid in seg.get("member_ids", [])
        )
        original_corrected = original_corrected or seg.get(
            "corrected_text", seg["raw_text"]
        )
        if original_corrected != seg["raw_text"]:
            channel(
                "Терминология ORIGINAL",
                original_corrected.strip(),
                "1565C0",
                seg["language"] == "he",
            )
        for target, record, color in (("ru", seg, "00A6D6"), ("en", en, "D88A00")):
            correction = record.get("translation_corrected_" + target)
            if correction and correction != record.get("translation_" + target):
                channel("Терминология " + target.upper(), correction, color)
        if group_needs_review(seg, segments_by_id):
            note = "⚠ Требуется проверка"
            blocks.append(note)
            p = paragraph(note)
            p.runs[0].font.size = Pt(8)
            p.runs[0].font.color.rgb = RGBColor.from_string("B35A16")
            group_paragraphs.append(p)
        # Keep each ordinary block on one page; Word can split oversized blocks.
        for p in group_paragraphs:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = Pt(2)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in {
            "val": "single",
            "sz": "4",
            "color": "D9E2EC",
            "space": "3",
        }.items():
            bottom.set(qn("w:" + key), value)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)
        blocks.append("─" * 44)
    atomic_write_text(paths[1], "\n".join(blocks) + "\n")
    tmp = out / ("docx-" + uuid.uuid4().hex + ".docx")
    try:
        doc.save(tmp)
        os.replace(tmp, paths[2])
    finally:
        tmp.unlink(missing_ok=True)
    return paths


def publish_exports(data: dict, config: dict, out: Path) -> list[Path]:
    """Stage all formats, lock the flat destination, restore prior files on failure."""
    key = text_hash(str(out.resolve()) + Path(data["source"]["path"]).stem)[:24]
    technical = ROOT / "temp/v156-publication" / key
    with run_lock(technical):
        stage = technical / uuid.uuid4().hex
        stage.mkdir(parents=True)
        staged = export_all(data, config, stage)
        out.mkdir(parents=True, exist_ok=True)
        targets = [out / path.name for path in staged]
        backups = {}
        for target in targets:
            if target.exists():
                backup = stage / (target.name + ".previous")
                shutil.copy2(target, backup)
                backups[target] = backup
        replaced = []
        try:
            for source, target in zip(staged, targets, strict=True):
                os.replace(source, target)
                replaced.append(target)
        except OSError as error:
            for target in reversed(replaced):
                if target in backups:
                    os.replace(backups[target], target)
                else:
                    target.unlink(missing_ok=True)
            raise PipelineError(
                "Не удалось заменить output; прежние файлы восстановлены. Закройте DOCX в Word и повторите запуск."
            ) from error
        for backup in backups.values():
            backup.unlink(missing_ok=True)
        stage.rmdir()
        return targets


def release_models() -> None:
    gc.collect()
    import torch

    if torch.cuda.is_available():
        torch.cuda.empty_cache()


def process(
    source: Path,
    config: dict,
    *,
    limit_seconds: float | None = None,
    fresh: bool = False,
    retry_translation: bool = False,
) -> tuple[dict, list[Path]]:
    started = time.perf_counter()
    # Offline is deliberate: no hidden downloads, redirects are delegated to HF only
    # during a separately managed model provisioning step.
    os.environ["HF_HUB_OFFLINE"] = "1"
    os.environ["HF_HUB_DISABLE_TELEMETRY"] = "1"
    os.environ["PYANNOTE_METRICS_ENABLED"] = "0"
    runtime_cache = inside_root(config["TEMP"]) / "runtime"
    runtime_cache.mkdir(parents=True, exist_ok=True)
    os.environ["MPLCONFIGDIR"] = str(runtime_cache / "matplotlib")
    os.environ["NUMBA_CACHE_DIR"] = str(runtime_cache / "numba")
    if config["FORCE_LANGUAGE"] == "zu":
        raise PipelineError(
            "Zulu ASR unsupported by current backend; NLLB zul_Latn supports trusted text only"
        )
    if not shutil.which("ffmpeg"):
        raise PipelineError("FFmpeg отсутствует в PATH")
    source = inside_root(source)
    identity = source_identity(source)
    runtime = runtime_info(config)
    paths = {
        "whisper": model_path(
            "Systran/faster-whisper-large-v3",
            config["WHISPER_PATH"],
            ("model.bin", "tokenizer.json", "config.json"),
        ),
        "translator": model_path(
            config["TRANSLATOR_MODEL"],
            config["TRANSLATOR_PATH"],
            ("config.json", "tokenizer.json", "pytorch_model.bin"),
        ),
        "diarization": model_path(
            "pyannote/speaker-diarization-3.1",
            config["DIARIZATION_PATH"],
            ("config.yaml",),
        ),
    }
    for dependency in ("segmentation-3.0", "wespeaker-voxceleb-resnet34-LM"):
        paths[dependency] = model_path(
            "pyannote/" + dependency, None, ("pytorch_model.bin",)
        )
    metadata = processing_metadata(config, runtime, paths)
    metadata["limit_seconds"] = limit_seconds
    base, fingerprint = cache_location(inside_root(config["TEMP"]), identity, metadata)
    if fresh:
        base = base.with_name(base.name + "-fresh-" + uuid.uuid4().hex[:8])
    data, folder = read_cache(base, identity, fingerprint)
    with run_lock(folder):
        if data is None and not fresh:
            data = migrate_ux_predecessor(
                inside_root(config["TEMP"]), identity, metadata
            )

        def checkpoint():
            if data is not None:
                sync_group_status(data)
                data["updated_at"] = utc_now()
                data["state_sha256"] = digest(
                    {k: v for k, v in data.items() if k != "state_sha256"}
                )
                validate_cache(data, identity, fingerprint)
                atomic_write_json(folder / "resume.json", data)

        audio = None
        try:
            if (
                data is None
                or not data["units"]
                or any(not u["asr_done"] for u in data["units"])
            ):
                wav = folder / "prepared.wav"
                if (
                    data is not None
                    and wav.exists()
                    and data.get("prepared_sha256") != file_hash(wav)
                ):
                    raise PipelineError(
                        "Prepared WAV изменён/повреждён; --fresh создаст новый namespace"
                    )
                if not wav.exists():
                    prepare_audio(source, wav)
                audio = load_audio(wav)
                if limit_seconds:
                    audio = audio[: int(limit_seconds * 16000)]
                if data is None:
                    duration = len(audio) / 16000
                    data = {
                        "schema_version": SCHEMA,
                        "pipeline_version": VERSION,
                        "source": identity,
                        "processing": metadata,
                        "processing_fingerprint": fingerprint,
                        "duration": duration,
                        "created_at": utc_now(),
                        "updated_at": utc_now(),
                        "diarization": [],
                        "exclusive_diarization": [],
                        "diarization_done": False,
                        "units": [],
                        "segments": [],
                        "export": {},
                        "metrics": {},
                        "warnings": [
                            "Zulu cannot be auto-identified by current Whisper language table"
                        ],
                    }
                data["prepared_sha256"] = file_hash(wav)
                if not data["diarization_done"]:
                    LOG.info("Диаризация 3.1, waveform API")
                    regular, exclusive = diarize(audio, config, paths["diarization"])
                    data.update(
                        diarization=regular,
                        exclusive_diarization=exclusive,
                        diarization_done=True,
                    )
                    checkpoint()
                backend = WhisperBackend(paths["whisper"], config, runtime)
                lid = TextLanguageValidator()
                try:
                    if not data["units"]:
                        data["units"] = language_units(audio, backend, config)
                        checkpoint()
                    for index, unit in enumerate(data["units"]):
                        if unit["asr_done"]:
                            continue
                        segments = transcribe_unit(
                            unit,
                            audio,
                            backend,
                            config,
                            data["diarization"],
                            data["exclusive_diarization"],
                        )
                        reconcile_unit(
                            unit,
                            segments,
                            data["units"],
                            index,
                            audio,
                            backend,
                            lid,
                            config,
                        )
                        data["segments"].extend(segments)
                        unit["segment_ids"] = [s["segment_id"] for s in segments]
                        unit["raw_pieces_sha256"] = digest(unit["raw_pieces"])
                        unit["asr_done"] = True
                        checkpoint()
                        LOG.info(
                            "ASR %d/%d | %s | %s",
                            index + 1,
                            len(data["units"]),
                            format_time(unit["start"]),
                            unit["language"],
                        )
                finally:
                    del backend
                    release_models()
            if not data.get("translation_groups"):
                data["translation_groups"] = build_translation_groups(data["segments"])
                checkpoint()
            import translation_v156 as mt

            mt.prepare_routing(sys.modules[__name__], data, config)
            translator = None
            try:
                for group in data["translation_groups"]:
                    group.setdefault(
                        "english",
                        {
                            "target_language": "en",
                            "translation_context": copy.deepcopy(
                                group.get("translation_context", {})
                            ),
                            "language": group["language"],
                            "raw_text": group["raw_text"],
                            "raw_sha256": group["raw_sha256"],
                            "translation_en": None,
                            "translation_status": "pending",
                            "chunks": [],
                            "warnings": [],
                        },
                    )
                records = [(group, "ru") for group in data["translation_groups"]]
                records += [
                    (group["english"], "en") for group in data["translation_groups"]
                ]
                for index, (seg, target_language) in enumerate(records):
                    eligible = (
                        seg["translation_status"] == "pending"
                        or any(pending_chunk(ch) for ch in seg["chunks"])
                        or retry_translation
                        and seg["translation_status"] in {"failed", "partial"}
                    )
                    if not eligible:
                        continue
                    if (
                        seg["language"] in NLLB
                        and seg["language"] != target_language
                        and seg["raw_text"].strip()
                        and config["TRANSLATE_MODE"] != "off"
                    ):
                        if translator is None:
                            translator = Translator(paths["translator"], config)
                    if seg["language"] in {"af", "nl"}:
                        mt.translate_group(
                            sys.modules[__name__],
                            seg,
                            translator,
                            config,
                            checkpoint,
                            retry_translation,
                            target_language,
                        )
                    else:
                        translate_segment(
                            seg,
                            translator,
                            config,
                            checkpoint,
                            retry_translation,
                            target_language,
                        )
                    LOG.info(
                        "NLLB %d/%d | %s | target=%s",
                        index + 1,
                        len(records),
                        seg["translation_status"],
                        target_language,
                    )
            finally:
                if translator is not None:
                    del translator
                    release_models()
            if source_identity(source) != identity:
                raise PipelineError("Source изменился во время обработки")
            elapsed = time.perf_counter() - started
            data["metrics"] = {
                "last_run_wall_seconds": elapsed,
                "last_run_rtf": elapsed / data["duration"],
                "note": "through translation, excluding final export; resume is last run only; torch VRAM excludes CTranslate2",
                "runtime": runtime,
            }
            try:
                import psutil

                memory = psutil.Process().memory_info()
                data["metrics"]["process_peak_working_set_mb"] = (
                    getattr(memory, "peak_wset", memory.rss) // 1024**2
                )
            except (ImportError, OSError):
                data["metrics"]["process_peak_working_set_mb"] = None
            import torch

            if torch.cuda.is_available():
                data["metrics"]["torch_peak_allocated_mb"] = (
                    torch.cuda.max_memory_allocated() // 1024**2
                )
            out = inside_root(config["OUTPUT"])
            exported = publish_exports(data, config, out)
            data["export"] = {
                "status": "complete",
                "files": [str(p) for p in exported],
                "at": utc_now(),
            }
            checkpoint()
            LOG.info(
                "RESULT %s | %.1fs | RTF %.3f | %s",
                exported_status(data),
                elapsed,
                elapsed / data["duration"],
                out,
            )
            for path in exported:
                LOG.info("OUTPUT %s", path)
            return data, exported
        finally:
            checkpoint()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=f"Audio2Text {VERSION}")
    parser.add_argument("source", nargs="?", type=Path)
    parser.add_argument("--config", type=Path, default=ROOT / "config_v156.json")
    parser.add_argument("--device", choices=["cpu", "cuda", "auto"])
    parser.add_argument("--force-language")
    parser.add_argument("--min-speakers", type=int)
    parser.add_argument("--max-speakers", type=int)
    parser.add_argument(
        "--limit-seconds",
        type=float,
        help="isolated cache and _SMOKE_PRO output; does not overwrite full run",
    )
    parser.add_argument("--fresh", action="store_true")
    parser.add_argument("--retry-translation", action="store_true")
    parser.add_argument("--validate-config", action="store_true")
    args = parser.parse_args(argv)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%H:%M:%S",
    )
    try:
        c = load_config(args.config)
        for arg, key in [
            ("device", "DEVICE"),
            ("force_language", "FORCE_LANGUAGE"),
            ("min_speakers", "MIN_SPEAKERS"),
            ("max_speakers", "MAX_SPEAKERS"),
        ]:
            if getattr(args, arg) is not None:
                c[key] = getattr(args, arg)
        c = validate_config(c)
        if args.validate_config:
            LOG.info("CONFIG PASS")
            return 0
        if args.limit_seconds is not None and (
            not math.isfinite(args.limit_seconds) or args.limit_seconds <= 0
        ):
            raise PipelineError("--limit-seconds должен быть положительным")
        source = args.source
        if source is None:
            files = sorted(
                p
                for p in inside_root(c["INPUT"]).iterdir()
                if p.suffix.lower() in c["SUPPORTED_EXTENSIONS"] and p.is_file()
            )
            if not files:
                raise PipelineError("Нет поддерживаемых файлов в input")
            if len(files) > 1:
                LOG.warning(
                    "Несколько файлов: выбран первый по имени. Для другого передайте полный путь."
                )
            source = files[0]
        data, _ = process(
            source,
            c,
            limit_seconds=args.limit_seconds,
            fresh=args.fresh,
            retry_translation=args.retry_translation,
        )
        return 2 if exported_status(data) == "partial" else 0
    except KeyboardInterrupt:
        LOG.warning(
            "Остановлено. Успешные стадии сохранены последним atomic checkpoint; повторите ту же команду."
        )
        return 130
    except Exception as error:
        # Never echo arbitrary HF request URLs/headers or credential-bearing exceptions.
        LOG.error(
            "%s",
            str(error)
            if isinstance(error, PipelineError)
            else f"{type(error).__name__}; обработка остановлена, проверьте окружение/сохранённые стадии",
        )
        return 1


if __name__ == "__main__":
    sys.exit(main())

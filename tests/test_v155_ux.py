"""Bilingual semantic reading view and pinned cache upgrade contracts."""

import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
import audio_pipeline_cuda_v155_FIXED as p
from test_v155 import cache, config, segment, TokenTranslator
from test_v155_final import adjacent


class EnglishTranslator(TokenTranslator):
    def generate(self, text):
        result = super().generate(text)
        result["text"] = "We will meet tomorrow."
        result["target_language"] = self.target_language
        return result


def english_record(group):
    return {
        "raw_text": group["raw_text"],
        "raw_sha256": group["raw_sha256"],
        "language": group["language"],
        "target_language": "en",
        "translation_status": "pending",
        "translation_en": None,
        "chunks": [],
        "warnings": [],
    }


class UXTests(unittest.TestCase):
    def test_pinned_upgrade_rejects_unreviewed_target_code(self):
        with patch.object(p, "UX_UPGRADE_TARGET_AST", "different"):
            self.assertIsNone(
                p.migrate_ux_predecessor(p.ROOT / "tests/.runtime", {}, {})
            )

    def test_explicit_semantic_review_is_visible_without_debug_details(self):
        g = segment("Test", "en")
        g["warnings"] = ["semantic_review_required"]
        self.assertTrue(p.group_needs_review(g, {}))

    def test_failed_english_makes_export_partial(self):
        d = cache()
        d["segments"][0]["translation_status"] = "needs_review"
        d["translation_groups"] = [{"english": {"translation_status": "failed"}}]
        self.assertEqual(p.exported_status(d), "partial")

    def test_english_inference_all_non_english_sources_and_resume(self):
        for lang in p.NLLB:
            if lang == "en":
                continue
            with self.subTest(language=lang):
                record = english_record(segment("A short test.", lang))
                translator = EnglishTranslator()
                p.translate_segment(record, translator, config(), target_language="en")
                calls = translator.calls
                self.assertGreater(calls, 0)
                self.assertEqual(translator.source, p.NLLB[lang])
                self.assertEqual(record["translation_en"], "We will meet tomorrow.")
                self.assertNotIn("translation_ru", record["chunks"][0])
                p.translate_segment(record, translator, config(), target_language="en")
                self.assertEqual(translator.calls, calls)
                p.validate_chunks(record["raw_text"], record["chunks"], "en")

    def test_source_english_bypass(self):
        record = english_record(segment())
        p.translate_segment(record, None, config(), target_language="en")
        self.assertEqual(record["translation_status"], "not_required")
        self.assertIsNone(record["translation_en"])
        self.assertFalse(record["chunks"])

    def test_english_latin_is_valid_but_russian_wrong_script_rejected(self):
        result = {
            "text": "We will meet together tomorrow",
            "generated_ids": [],
            "target_token_count": 8,
        }
        self.assertFalse(p.generation_defects("Ons sal môre ontmoet", result, 8, "en"))
        result["text"] = "Мы завтра все вместе встретимся"
        self.assertIn(
            "wrong_target_script",
            p.generation_defects("Ons sal môre ontmoet", result, 8, "en"),
        )

    def test_semantic_view_order_colors_no_orphans_no_debug(self):
        d = cache()
        d["segments"] = adjacent([" Eerste deel", " tweede deel."])
        d["translation_groups"] = p.build_translation_groups(d["segments"])
        group = d["translation_groups"][0]
        group.update(translation_ru="Полная фраза.", translation_status="needs_review")
        group["english"] = english_record(group)
        group["english"].update(
            translation_en="The entire sentence.", translation_status="needs_review"
        )
        for seg in d["segments"]:
            seg["diagnostics"] = ["vad_alternative_differs", "generated_ids"]
        base = p.ROOT / "tests/.runtime"
        base.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=base) as tmp:
            files = p.export_all(d, config(), Path(tmp))
            txt = files[1].read_text("utf-8")
            self.assertEqual(txt.count("[ORIGINAL/AF]"), 1)
            self.assertEqual(txt.count("[00:00:00–00:00:03] A | AF"), 1)
            self.assertIn(
                "[ORIGINAL/AF]\nEerste deel tweede deel.\n[RU]\nПолная фраза.\n[EN]\nThe entire sentence.",
                txt,
            )
            doc = Document(files[2])
            paragraphs = [x.text for x in doc.paragraphs]
            start = paragraphs.index("ORIGINAL · AF")
            self.assertEqual(
                paragraphs[start : start + 6],
                [
                    "ORIGINAL · AF",
                    "Eerste deel tweede deel.",
                    "RU",
                    "Полная фраза.",
                    "EN",
                    "The entire sentence.",
                ],
            )
            for offset, color in enumerate(
                ["1565C0", "1565C0", "00A6D6", "00A6D6", "D88A00", "D88A00"]
            ):
                par = doc.paragraphs[start + offset]
                self.assertTrue(par.paragraph_format.keep_with_next)
                self.assertTrue(all(str(r.font.color.rgb) == color for r in par.runs))
            for content in [txt, "\n".join(paragraphs)]:
                for forbidden in [
                    "vad_",
                    "generated_ids",
                    "needs_review",
                    "t00000",
                    "confidence",
                    "NLLB",
                    "семантически",
                ]:
                    self.assertNotIn(forbidden, content)
                self.assertNotIn("⚠", content)
            self.assertEqual(len(p.read_json(files[0])["segments"]), 2)

    def test_ru_and_en_source_render_no_duplicate_translation(self):
        base = p.ROOT / "tests/.runtime"
        for language in ["ru", "en"]:
            d = cache()
            d["segments"] = [
                segment("Исходник" if language == "ru" else "Original", language)
            ]
            d["translation_groups"] = p.build_translation_groups(d["segments"])
            g = d["translation_groups"][0]
            g.update(
                translation_ru="Русский перевод", translation_status="needs_review"
            )
            g["english"] = english_record(g)
            g["english"].update(
                translation_en="English translation", translation_status="needs_review"
            )
            with tempfile.TemporaryDirectory(dir=base) as tmp:
                files = p.export_all(d, config(), Path(tmp))
                txt = files[1].read_text("utf-8")
                self.assertNotIn("[" + language.upper() + "]", txt)
                self.assertIn("[" + ("EN" if language == "ru" else "RU") + "]", txt)

    def test_english_cache_tamper_and_status(self):
        ss = adjacent([" Eerste deel", " tweede deel."])
        gg = p.build_translation_groups(ss)
        gg[0]["english"] = english_record(gg[0])
        d = {"segments": ss, "translation_groups": gg}
        p.validate_translation_groups(d)
        gg[0]["english"]["raw_text"] += " altered"
        with self.assertRaises(ValueError):
            p.validate_translation_groups(d)

    def test_pinned_migration_preserves_data_and_rejects_changed_inputs(self):
        current = {
            "algorithm_sha256": "new",
            "build": p.BUILD,
            "mapping_version": 3,
            "translation_channels": ["ru", "en"],
            "config": {"BEAM_SIZE": 5},
            "models": {"whisper": "unchanged"},
        }
        old = copy.deepcopy(current)
        old.pop("translation_channels")
        old.update(
            algorithm_sha256=p.UX_PREDECESSOR_SHA256,
            build="v155-reconciliation-groups-tokenizer-4",
            mapping_version=2,
        )
        d = cache()
        d["translation_groups"] = p.build_translation_groups(d["segments"])
        d["translation_groups"][0].update(
            translation_ru="Перевод", translation_status="needs_review"
        )
        d.update(processing=old, processing_fingerprint=p.digest(old))
        d["state_sha256"] = p.digest(
            {k: v for k, v in d.items() if k != "state_sha256"}
        )
        base = p.ROOT / "tests/.runtime"
        with tempfile.TemporaryDirectory(dir=base) as tmp:
            folder, _ = p.cache_location(Path(tmp), d["source"], old)
            p.atomic_write_json(folder / "resume.json", d)
            before = (folder / "resume.json").read_bytes()
            migrated = p.migrate_ux_predecessor(Path(tmp), d["source"], current)
            self.assertEqual(migrated["segments"], d["segments"])
            self.assertEqual(migrated["translation_groups"], d["translation_groups"])
            self.assertEqual(migrated["processing_fingerprint"], p.digest(current))
            for changed in ["config", "models"]:
                wrong = copy.deepcopy(current)
                wrong[changed] = {"changed": True}
                self.assertIsNone(
                    p.migrate_ux_predecessor(Path(tmp), d["source"], wrong)
                )
            wrong_source = {**d["source"], "mtime_ns": 21}
            self.assertIsNone(
                p.migrate_ux_predecessor(Path(tmp), wrong_source, current)
            )
            self.assertEqual((folder / "resume.json").read_bytes(), before)

    def test_english_resume_after_interruption_keeps_successful_chunk(self):
        record = english_record(segment("First sentence. Second sentence.", "af"))

        class Interrupted(EnglishTranslator):
            def generate(self, text):
                if self.calls == 1:
                    raise KeyboardInterrupt
                return super().generate(text)

        t = Interrupted()
        with self.assertRaises(KeyboardInterrupt):
            p.translate_segment(record, t, config(), target_language="en")
        saved = copy.deepcopy(record["chunks"][0])
        t = EnglishTranslator()
        p.translate_segment(record, t, config(), target_language="en")
        self.assertEqual(t.calls, 1)
        self.assertEqual(record["chunks"][0], saved)

    def test_whisper_and_diarization_core_unchanged(self):
        import ast

        old = ast.parse(
            (
                p.ROOT / "backups/v155_before_ux/audio_pipeline_cuda_v155_FIXED.py"
            ).read_text("utf-8")
        )
        new = ast.parse(Path(p.__file__).read_text("utf-8"))
        names = {
            "WhisperBackend",
            "prepare_audio",
            "diarize",
            "normalize_diarization",
            "speaker_for",
            "language_units",
            "transcribe_unit",
            "reconcile_unit",
            "build_translation_groups",
        }

        def selected(tree):
            return {
                n.name: ast.dump(n, include_attributes=False)
                for n in tree.body
                if getattr(n, "name", None) in names
            }

        self.assertEqual(selected(old), selected(new))


if __name__ == "__main__":
    unittest.main()
